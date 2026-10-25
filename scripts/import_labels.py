#!/usr/bin/env python3
"""Turn the lawyer's filled labels (csv or xlsx) into examples/<document>.expected.json for `npm run eval`.

Usage: python3 scripts/import_labels.py examples/labels/labels_template.xlsx
Rows with an empty expected_status are skipped. Several acceptable statuses: "risk_found|insufficient_data".
"""
import csv, json, sys, collections, pathlib

src = pathlib.Path(sys.argv[1])
if src.suffix == ".xlsx":
    from openpyxl import load_workbook
    ws = load_workbook(src, read_only=True).active
    it = list(ws.iter_rows(values_only=True))
    head, rows = [str(h) for h in it[0]], [dict(zip(it[0], r)) for r in it[1:]]
else:
    rows = list(csv.DictReader(open(src, encoding="utf-8-sig")))
out = collections.defaultdict(dict)
notes = collections.defaultdict(dict)
for r in rows:
    doc, rid, st = (r.get("document") or "").strip(), (r.get("risk_id") or "").strip(), (r.get("expected_status") or "").strip()
    if not (doc and rid and st):
        continue
    parts = [p.strip() for p in st.split("|")]
    bad = [p for p in parts if p not in ("risk_found", "no_risk", "insufficient_data")]
    if bad:
        sys.exit(f"{doc} {rid}: unknown status {bad}")
    out[doc][rid] = parts[0] if len(parts) == 1 else parts
for doc, risks in out.items():
    p = pathlib.Path("examples") / f"{doc}.expected.json"
    p.write_text(json.dumps({"note": "Labelled by the lawyer.", "risks": risks}, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"{p}: {len(risks)} labels")
